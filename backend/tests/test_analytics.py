"""PostgreSQL reporting/extract regressions with isolated local output."""
from datetime import UTC,datetime,timedelta
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
import json
import pytest
from pydantic import ValidationError
from sqlalchemy import select,func
from openpyxl import load_workbook
from app import models as m
from app.core.config import settings
from app.schemas import analytics as s
from app.services.analytics import reports as r,extracts as e,data,seed
from app.services.core_hr import employment,records
from app.schemas import core_hr as hr
from app.services.core_hr.common import Conflict,InvalidOperation
from test_core_hr import db,refs,client,headers,request

@pytest.fixture(autouse=True)
def output_dir(tmp_path,monkeypatch):monkeypatch.setattr(settings,'extract_output_dir',tmp_path)

def report(**kwargs):return s.ReportCreate(**(dict(code=uuid4().hex[:24],name='Workforce Assignment Summary',domain='CORE_HR_WORKERS',selected_columns=['person_number','annual_base_salary'])|kwargs))
def extract(**kwargs):return s.ExtractCreate(**(dict(code=uuid4().hex[:24],name='Worker Changes Export',extract_type='WORKER_CHANGES',output_format='JSON')|kwargs))

@pytest.mark.parametrize('change',[
 {'domain':'SQL'},{'selected_columns':['password_hash']},{'selected_columns':['person_number','person_number']},
 {'filters':[{'field':'password_hash','operator':'equals','value':'x'}]},
 {'filters':[{'field':'person_number','operator':'greater_than','value':'x'}]},
 {'filters':[{'field':'annual_base_salary','operator':'equals','value':'NaN'}]},
 {'filters':[{'field':'start_date','operator':'date_on_or_after','value':'yesterday'}]},
 {'filters':[{'field':'person_number','operator':'in','value':'x'}]},
 {'sort_definition':[{'field':'DROP TABLE','direction':'asc'}]},
])
def test_reject_arbitrary_query(change):
 with pytest.raises(ValidationError):report(**change)


def test_filters_sort_pagination_snapshot_exports(db,refs):
 h=employment.hire(db,request(refs));d=r.create_report(db,report(filters=[s.Filter(field='person_number',operator='equals',value=h.person.person_number)],sort_definition=[s.Sort(field='annual_base_salary',direction='desc')]))
 run=r.run_report(db,d.id,s.ReportRequest(as_of='2024-01-01'));assert run.status=='COMPLETED' and run.row_count==1
 assert r.results(db,run.id)['rows'][0]['annual_base_salary']=='100000.25'
 assert r.results(db,run.id,1)['rows']==[]
 r.update_report(db,d.id,report(code=d.code,selected_columns=['employee_name']))
 assert r.results(db,run.id)['columns']==['person_number','annual_base_salary']
 csv=r.export(db,run.id,'csv').decode('utf-8-sig');assert csv.startswith('Person Number,Annual Base Salary')
 book=load_workbook(BytesIO(r.export(db,run.id,'xlsx')));assert book.active.cell(2,2).value=='100000.25';book.close()

@pytest.mark.parametrize('op,value,expected',[('equals','2',True),('not_equals','3',True),('in',['1','2'],True),('greater_than','1',True),('less_than','1',False)])
def test_typed_filters(op,value,expected):assert r.matches({'n':Decimal('2')},s.Filter(field='n',operator=op,value=value),'decimal') is expected

def test_formula_exports():
 values=[{'name':'  =SUM(A1)'},{'name':'@test'},{'name':'-cmd'},{'name':'+cmd'}]
 assert "'  =SUM(A1)" in r.csv_bytes(['name'],values).decode('utf-8-sig')
 book=load_workbook(BytesIO(r.xlsx_bytes(['name'],values)));assert book.active.cell(2,1).data_type=='s' and book.active.cell(2,1).value.startswith("'");book.close()


def test_full_incremental_and_failed_watermark(db,refs,monkeypatch):
 h=employment.hire(db,request(refs));d=e.create(db,extract(configuration=s.ExtractConfig(filters=[s.Filter(field='person_number',operator='equals',value=h.person.person_number)])))
 first=e.run(db,d.id,s.ExtractRequest(mode='FULL'));assert first.status=='COMPLETED' and first.row_count==1
 assert len(json.loads(e.download(db,first.id)[1].read_text()))==1
 assert e.run(db,d.id,s.ExtractRequest(mode='INCREMENTAL')).row_count==0
 person=db.get(m.Person,h.person.id);person.preferred_name='Sai Kiran';person.updated_at=datetime.now(UTC);db.flush()
 changed=e.run(db,d.id,s.ExtractRequest(mode='INCREMENTAL'));assert changed.row_count==1
 old=d.last_successful_run_at
 def fail(*a,**k):raise OSError('private path')
 monkeypatch.setattr(e,'write_output',fail)
 failed=e.run(db,d.id,s.ExtractRequest(mode='INCREMENTAL'));assert failed.status=='FAILED' and d.last_successful_run_at==old and 'private' not in failed.error_message
 assert not failed.output_filename
 with pytest.raises(Conflict):e.download(db,failed.id)


def test_exact_boundaries_and_paths():
 start=datetime.now(UTC);end=start+timedelta(seconds=1)
 assert not e.interval(start,start,end)
 assert e.interval(end,start,end)
 assert not e.interval(end+timedelta(microseconds=1),start,end)
 for path in ('../secret.csv','C:/secret.json','not-a-uuid.csv'):
  with pytest.raises(InvalidOperation):e.output_path(path)


def test_scope_freezes_and_disabled(db):
 d=e.create(db,extract());e.run(db,d.id,s.ExtractRequest())
 with pytest.raises(Conflict):e.update(db,d.id,extract(code=d.code,extract_type='PAYROLL_RESULTS'))
 payload=extract(code=d.code,is_active=False);e.update(db,d.id,payload)
 with pytest.raises(Conflict):e.run(db,d.id,s.ExtractRequest())
 rd=r.create_report(db,report(is_active=False))
 with pytest.raises(Conflict):r.run_report(db,rd.id,s.ReportRequest())


def test_all_domains_and_seed(db):
 for domain,columns in [('PAYROLL_RESULTS',['net_pay']),('FBP_ALLOCATIONS',['elected_amount']),('IMPORT_HISTORY',['filename'])]:
  rd=r.create_report(db,report(domain=domain,selected_columns=columns));assert r.run_report(db,rd.id,s.ReportRequest()).status=='COMPLETED'
 for kind in ('WORKER_SNAPSHOT','PAYROLL_RESULTS','FBP_ELECTIONS'):
  d=e.create(db,extract(extract_type=kind));assert e.run(db,d.id,s.ExtractRequest()).status=='COMPLETED'
 seed.seed(db);before=db.scalar(select(func.count()).select_from(m.ReportDefinition));assert seed.seed(db)==0;assert db.scalar(select(func.count()).select_from(m.ReportDefinition))==before


def test_safe_failure(db,monkeypatch):
 d=r.create_report(db,report())
 def fail(*args):raise RuntimeError('raw database secret')
 monkeypatch.setattr(data,'domain_rows',fail)
 run=r.run_report(db,d.id,s.ReportRequest());assert run.status=='FAILED' and 'secret' not in run.error_message

@pytest.mark.parametrize('role',['HR','ADMIN'])
def test_staff_apis(client,db,role):
 auth=headers(db,role);db.commit()
 assert client.get('/reports/metadata',headers=auth).status_code==200
 response=client.post('/reports',headers=auth,json=report().model_dump(mode='json'));assert response.status_code==201,response.text
 run=client.post('/reports/'+response.json()['id']+'/run',headers=auth,json={});assert run.status_code==200 and run.json()['status']=='COMPLETED',run.text
 rid=run.json()['id']
 for suffix in ('','/results','/export.csv','/export.xlsx'):assert client.get('/reports/runs/'+rid+suffix,headers=auth).status_code==200
 response=client.post('/extracts/definitions',headers=auth,json=extract().model_dump(mode='json'));assert response.status_code==201,response.text
 run=client.post('/extracts/definitions/'+response.json()['id']+'/run',headers=auth,json={'mode':'FULL'});assert run.status_code==200 and run.json()['status']=='COMPLETED',run.text
 assert client.get('/extracts/runs/'+run.json()['id']+'/download',headers=auth).status_code==200


def test_employee_denied(client,db):
 auth=headers(db,'EMPLOYEE');db.commit();uid=str(uuid4())
 for path in ('/reports','/reports/metadata','/reports/runs/'+uid+'/export.xlsx','/extracts/definitions','/extracts/runs/'+uid+'/download'):
  assert client.get(path,headers=auth).status_code==403
 for path in ('/reports/'+uid+'/run','/extracts/definitions/'+uid+'/run'):
  assert client.post(path,headers=auth,json={}).status_code==403
 assert client.get('/reports').status_code==401


def test_migration_upgrade_downgrade_upgrade_isolated_schema():
    import importlib.util
    from pathlib import Path
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import inspect, text
    from app.core.database import engine
    schema='analytics_test_'+uuid4().hex
    modules={}
    for path in Path('alembic/versions').glob('*.py'):
        spec=importlib.util.spec_from_file_location(path.stem,path)
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        modules[module.revision]=module
    ordered=[];head='57dcd64fceaa'
    while head:
        ordered.insert(0,modules[head]);head=modules[head].down_revision
    with engine.connect() as connection:
        transaction=connection.begin()
        try:
            connection.execute(text(f'CREATE SCHEMA "{schema}"'))
            connection.execute(text(f'SET LOCAL search_path TO "{schema}", public'))
            with Operations.context(MigrationContext.configure(connection)):
                for migration in ordered:migration.upgrade()
                assert 'report_runs' in inspect(connection).get_table_names(schema=schema)
                ordered[-1].downgrade()
                assert 'report_runs' not in inspect(connection).get_table_names(schema=schema)
                assert 'payroll_results' in inspect(connection).get_table_names(schema=schema)
                ordered[-1].upgrade()
                assert 'extract_runs' in inspect(connection).get_table_names(schema=schema)
        finally:transaction.rollback()


def test_orphan_cleanup_keeps_referenced_files(db):
 import os
 from scripts.cleanup_extract_orphans import cleanup
 d=e.create(db,extract());run=e.run(db,d.id,s.ExtractRequest())
 referenced=e.output_path(run.output_filename);orphan=e.output_path(str(uuid4())+'.csv');orphan.write_text('synthetic')
 past=(datetime.now(UTC)-timedelta(days=10)).timestamp()
 os.utime(referenced,(past,past));os.utime(orphan,(past,past))
 assert cleanup(db,7)==1 and referenced.exists() and not orphan.exists()


def test_compensation_only_change_detected(db,refs):
 h=employment.hire(db,request(refs));d=e.create(db,extract(configuration=s.ExtractConfig(filters=[s.Filter(field='person_number',operator='equals',value=h.person.person_number)])))
 e.run(db,d.id,s.ExtractRequest())
 changed=employment.change_compensation(db,h.assignment.id,hr.AssignmentCompensationCreate(effective_from='2026-01-01',annual_base_salary='200000.25',currency='INR'))
 # Fixture shares one outer transaction, so explicitly simulate a later transaction timestamp.
 changed.updated_at=datetime.now(UTC);db.flush()
 assert e.run(db,d.id,s.ExtractRequest(mode='INCREMENTAL')).row_count==1
